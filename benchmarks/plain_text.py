"""Plain text extraction, fixed-size chunking and BM25 retrieval for the conventional baselines.

This is deliberately the textbook approach: pypdf for PDF pages, python-docx for Word paragraphs
and tables, openpyxl for spreadsheet rows. No layout analysis, no OCR, no embeddings.
"""

import math
import re
from collections import Counter
from pathlib import Path

SUPPORTED = {".pdf", ".docx", ".xlsx"}


def extract_units(folder) -> list:
    """Return text units in a fixed order: [{"file", "page", "text"}].

    PDF units are pages (1-based page numbers); a Word file is one unit and a workbook gives one
    unit per sheet, both with page None because they have no page geometry.
    """
    folder = Path(folder)
    order = {".pdf": 0, ".docx": 1, ".xlsx": 2}
    files = sorted((p for p in folder.iterdir() if p.suffix.lower() in SUPPORTED),
                   key=lambda p: (order[p.suffix.lower()], p.name))
    units = []
    for path in files:
        ext = path.suffix.lower()
        if ext == ".pdf":
            from pypdf import PdfReader

            for i, page in enumerate(PdfReader(str(path)).pages, start=1):
                units.append({"file": path.name, "page": i, "text": page.extract_text() or ""})
        elif ext == ".docx":
            units.append({"file": path.name, "page": None, "text": docx_text(path)})
        else:
            from openpyxl import load_workbook

            wb = load_workbook(str(path), data_only=True, read_only=True)
            for ws in wb.worksheets:
                rows = []
                for row in ws.iter_rows(values_only=True):
                    cells = [str(v) for v in row if v is not None and str(v).strip()]
                    if cells:
                        rows.append(" | ".join(cells))
                units.append({"file": path.name, "page": None, "sheet": ws.title,
                              "text": f"Sheet: {ws.title}\n" + "\n".join(rows)})
            wb.close()
    return units


def docx_text(path) -> str:
    """Paragraphs and table rows of a .docx file in document order."""
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = docx.Document(str(path))
    lines = []
    for child in d.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = Paragraph(child, d).text
            if text.strip():
                lines.append(text)
        elif tag == "tbl":
            for row in Table(child, d).rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    lines.append(" | ".join(cells))
    return "\n".join(lines)


def chunk_units(units, size: int = 1000, overlap: int = 200) -> list:
    """Fixed-size character windows inside each unit (chunks never cross a page or file)."""
    if overlap >= size:
        raise ValueError("overlap must be smaller than size")
    chunks = []
    for u in units:
        text = re.sub(r"[ \t]+", " ", u["text"]).strip()
        if not text:
            continue
        start = 0
        while True:
            piece = text[start:start + size]
            chunks.append({"file": u["file"], "page": u["page"], "text": piece})
            if start + size >= len(text):
                break
            start += size - overlap
    return chunks


def tokenize(text: str) -> list:
    return re.findall(r"[a-z0-9]+", text.lower())


class BM25:
    """Okapi BM25 over a list of documents (k1 = 1.5, b = 0.75)."""

    def __init__(self, docs, k1: float = 1.5, b: float = 0.75):
        self.k1, self.b = k1, b
        self.docs = [tokenize(d) for d in docs]
        self.tf = [Counter(d) for d in self.docs]
        self.avgdl = sum(len(d) for d in self.docs) / max(len(self.docs), 1)
        df = Counter(t for d in self.docs for t in set(d))
        n = len(self.docs)
        self.idf = {t: math.log(1 + (n - f + 0.5) / (f + 0.5)) for t, f in df.items()}

    def scores(self, query: str) -> list:
        q = tokenize(query)
        out = []
        for tf, doc in zip(self.tf, self.docs):
            s = 0.0
            for t in q:
                if t not in tf:
                    continue
                f = tf[t]
                s += self.idf[t] * f * (self.k1 + 1) / (f + self.k1 * (1 - self.b + self.b * len(doc) / self.avgdl))
            out.append(s)
        return out

    def top_k(self, query: str, k: int = 5) -> list:
        s = self.scores(query)
        return sorted(range(len(s)), key=lambda i: (-s[i], i))[:k]
