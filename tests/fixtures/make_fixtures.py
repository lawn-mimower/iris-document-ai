#!/usr/bin/env python3
"""
Builds the small synthetic test fixtures used by the test-suite and the
example commands. Every company, person and value here is fictional.

    python tests/fixtures/make_fixtures.py

Outputs (next to this script):
    form_empty.pdf / form_empty.png   blank demo form (labels, boxes, checkboxes)
    form_filled.pdf                   the same form filled in ("golden sample")
    form_new_data.json                values to place on the blank form
    dam_docs/                         fictional financial statements (PDF, XLSX, DOCX)
    dam_ground_truth.json             questions + expected answers for dam_docs/
"""

import json
from pathlib import Path

HERE = Path(__file__).resolve().parent

# --------------------------------------------------------------------------
# Demo form
# --------------------------------------------------------------------------

PAGE_W, PAGE_H = 612, 792  # US Letter, points
CHECKBOX = 4.5             # side of a checkbox in points
BOX_H = 20                 # height of a text input box in points

# (label, x_label, y, box_x, box_w)
TEXT_FIELDS = [
    ("1. Full Name:", 60, 690, 200, 170),
    ("2. Date of Birth (DD/MM/YYYY):", 60, 650, 240, 120),
    ("3. Email Address:", 60, 610, 200, 170),
    ("6. Membership Number:", 60, 490, 200, 110),
]

# (label, y, [(option, x_box)])
CHECK_GROUPS = [
    ("4. Gender:", 570, [("Male", 200), ("Female", 270), ("Other", 350)]),
    ("5. Relationship to Primary Member:", 530, [("Self", 260), ("Spouse", 320), ("Child", 400)]),
]

FILLED_VALUES = {
    "1. Full Name:": "Alex Example",
    "2. Date of Birth (DD/MM/YYYY):": "14/07/1988",
    "3. Email Address:": "alex@example.invalid",
    "6. Membership Number:": "DEMO-0042",
}
FILLED_CHECKS = {"4. Gender:": "Female", "5. Relationship to Primary Member:": "Self"}

NEW_DATA = {
    "Full Name": "Jordan Sample",
    "Date of Birth": "02/11/1990",
    "Email Address": "jordan@example.invalid",
    "Gender": "Male",
    "Relationship to Primary Member": "Spouse",
    "Membership Number": "DEMO-0107",
}


def draw_form(path: Path, filled: bool) -> None:
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=(PAGE_W, PAGE_H), invariant=1)
    c.setTitle("Demo membership form (fictional)")
    c.setFont("Helvetica-Bold", 14)
    c.drawString(60, 750, "DEMO MEMBERSHIP APPLICATION")
    c.setFont("Helvetica", 8)
    c.drawString(60, 736, "Fictional sample form for testing - Example Rowing Club does not exist.")

    c.setLineWidth(0.75)
    for label, x_label, y, box_x, box_w in TEXT_FIELDS:
        c.setFont("Helvetica", 10)
        c.drawString(x_label, y + 6, label)
        c.rect(box_x, y, box_w, BOX_H, stroke=1, fill=0)
        if filled:
            c.setFont("Helvetica", 10)
            c.drawString(box_x + 4, y + 6, FILLED_VALUES[label])

    c.setLineWidth(0.5)
    for label, y, options in CHECK_GROUPS:
        c.setFont("Helvetica", 10)
        c.drawString(60, y, label)
        for option, x_box in options:
            c.rect(x_box, y, CHECKBOX, CHECKBOX, stroke=1, fill=0)
            c.drawString(x_box + CHECKBOX + 4, y, option)
            if filled and FILLED_CHECKS[label] == option:
                c.setFont("Helvetica-Bold", 6)
                c.drawString(x_box + 0.8, y + 0.6, "x")
                c.setFont("Helvetica", 10)

    # Signature / date lines (underline style fields)
    c.setLineWidth(0.75)
    c.setFont("Helvetica", 10)
    c.drawString(60, 440, "Signature:")
    c.line(120, 438, 300, 438)
    c.drawString(340, 440, "Date:")
    c.line(372, 438, 480, 438)
    c.save()


def render_png(pdf_path: Path, png_path: Path, dpi: int = 150) -> None:
    import fitz

    with fitz.open(pdf_path) as doc:
        pix = doc[0].get_pixmap(dpi=dpi, colorspace=fitz.csGRAY)
        pix.save(str(png_path))


# --------------------------------------------------------------------------
# Fictional financial statements for the DAM pipeline
# --------------------------------------------------------------------------

COMPANY = {
    "name": "Examplecorp Widgets Private Limited",
    "cin": "U00000XX0000PTC000000",
    "address": "Unit 7, Imaginary Industrial Park, Sampletown, Nowhere State 000000",
    "email": "accounts@examplecorp.invalid",
    "fy_from": "01/04/2023",
    "fy_to": "31/03/2024",
    "board_date": "05/09/2024",
    "agm_date": "28/09/2024",
    "authorised_capital": "Rs. 10,00,000",
    "paid_up_capital": "Rs. 5,00,000",
    "revenue": "Rs. 1,23,45,000",
    "pat": "Rs. 8,76,000",
    "industry": "Manufacturing (widgets and fasteners)",
}


def make_financial_pdf(path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table

    s = getSampleStyleSheet()
    c = COMPANY
    story = [
        Paragraph(f"{c['name']} (fictional)", s["Title"]),
        Paragraph("Financial Statements for the year ended 31 March 2024 - synthetic test document", s["Normal"]),
        Spacer(1, 12),
        Paragraph("Corporate Information", s["Heading2"]),
        Paragraph(f"Corporate Identity Number (CIN): {c['cin']}", s["Normal"]),
        Paragraph(f"Registered office: {c['address']}", s["Normal"]),
        Paragraph(f"E-mail: {c['email']}", s["Normal"]),
        Paragraph(f"Financial year: {c['fy_from']} to {c['fy_to']}", s["Normal"]),
        Paragraph(f"Principal activity: {c['industry']}", s["Normal"]),
        Spacer(1, 12),
        Paragraph("Share Capital", s["Heading2"]),
        Paragraph(
            f"The authorised share capital of the company is {c['authorised_capital']} divided into "
            f"1,00,000 equity shares of Rs. 10 each. The issued, subscribed and paid-up capital is "
            f"{c['paid_up_capital']}.",
            s["Normal"],
        ),
        Spacer(1, 12),
        Paragraph("Statement of Profit and Loss (summary)", s["Heading2"]),
        Table(
            [
                ["Particulars", "FY 2023-24", "FY 2022-23"],
                ["Revenue from operations", "1,23,45,000", "1,01,20,000"],
                ["Other income", "1,10,000", "95,000"],
                ["Total expenses", "1,12,60,000", "95,40,000"],
                ["Profit after tax", "8,76,000", "4,90,000"],
            ]
        ),
        Spacer(1, 12),
        Paragraph("Approval of Financial Statements", s["Heading2"]),
        Paragraph(
            f"These financial statements were approved by the Board of Directors at its meeting held on "
            f"{c['board_date']}. The Annual General Meeting was held on {c['agm_date']}.",
            s["Normal"],
        ),
        Paragraph(
            "Secretarial audit under section 204 is not applicable to the company. "
            "The company has no subsidiaries, so consolidated financial statements are not required.",
            s["Normal"],
        ),
    ]
    SimpleDocTemplate(str(path), pagesize=A4, title="Fictional financial statements", invariant=1).build(story)


def make_trial_balance_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Trial Balance"
    ws.append(["Account", "Debit (Rs.)", "Credit (Rs.)"])
    for row in [
        ["Equity share capital (Examplecorp, fictional)", 0, 500000],
        ["Reserves and surplus", 0, 1876000],
        ["Trade receivables", 1450000, 0],
        ["Inventories", 980000, 0],
        ["Cash and bank balances", 620000, 0],
        ["Trade payables", 0, 674000],
    ]:
        ws.append(row)
    wb.save(path)


def make_directors_report_docx(path: Path) -> None:
    import docx

    d = docx.Document()
    d.core_properties.author = "Example author"
    d.core_properties.comments = "Fictional sample document for tests"
    d.add_heading("Directors' Report (fictional sample)", level=1)
    d.add_paragraph(
        f"The Directors present the annual report of {COMPANY['name']} for the financial year "
        f"{COMPANY['fy_from']} to {COMPANY['fy_to']}."
    )
    d.add_heading("Financial highlights", level=2)
    d.add_paragraph(
        f"Revenue from operations increased to {COMPANY['revenue']} and profit after tax was {COMPANY['pat']}."
    )
    d.add_heading("Board meetings", level=2)
    d.add_paragraph(
        f"The financial statements were approved at the board meeting held on {COMPANY['board_date']}. "
        "Directors: A. Placeholder (Managing Director) and B. Sample (Director)."
    )
    d.save(path)


GROUND_TRUTH = {
    "description": "Synthetic ground truth for tests/fixtures/dam_docs (fictional company).",
    "questions": [
        {"id": "Q1", "question": "Corporate Identity Number (CIN) of company", "answer": COMPANY["cin"]},
        {"id": "Q2", "question": "Name of the company", "answer": COMPANY["name"]},
        {"id": "Q3", "question": "Authorised capital of the company (in Rs.)", "answer": "10,00,000"},
        {"id": "Q4", "question": "Date of Board of Directors' meeting in which financial statements are approved", "answer": COMPANY["board_date"]},
        {"id": "Q5", "question": "Whether Secretarial Audit is applicable (Yes/No)", "answer": "No"},
    ],
}


def main() -> None:
    draw_form(HERE / "form_empty.pdf", filled=False)
    draw_form(HERE / "form_filled.pdf", filled=True)
    render_png(HERE / "form_empty.pdf", HERE / "form_empty.png")
    (HERE / "form_new_data.json").write_text(json.dumps(NEW_DATA, indent=2) + "\n")

    dam = HERE / "dam_docs"
    dam.mkdir(exist_ok=True)
    make_financial_pdf(dam / "examplecorp_financial_statements.pdf")
    make_trial_balance_xlsx(dam / "examplecorp_trial_balance.xlsx")
    make_directors_report_docx(dam / "examplecorp_directors_report.docx")
    (HERE / "dam_ground_truth.json").write_text(json.dumps(GROUND_TRUTH, indent=2) + "\n")

    for p in sorted(HERE.rglob("*")):
        if p.is_file() and p.suffix != ".py":
            print(f"{p.relative_to(HERE)}: {p.stat().st_size} bytes")


if __name__ == "__main__":
    main()
